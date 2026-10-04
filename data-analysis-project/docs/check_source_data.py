"""Download the cited public source and record a small, reproducible feasibility check."""
from pathlib import Path
import hashlib
import json
import platform
from datetime import datetime, timezone
import requests
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data' / 'raw'
RAW.mkdir(parents=True, exist_ok=True)
API = 'https://data.mendeley.com/public-api/datasets/386vmj2tbk/files?folder_id=root&version=4&$start=0&$limit=1000'
response = requests.get(API, timeout=60)
response.raise_for_status()
file_list = response.json()
selected = {
    'Dataset of health insurance portfolio.xlsx': 'health_insurance_portfolio.xlsx',
    'Descriptive of the variables.xlsx': 'variable_dictionary.xlsx',
}
manifest = {'dataset': 'Dataset of health insurance portfolio',
            'version': 4, 'doi': '10.17632/386vmj2tbk.4',
            'source': 'https://data.mendeley.com/datasets/386vmj2tbk/4',
            'license': 'CC BY 4.0',
            'checked_at_utc': datetime.now(timezone.utc).isoformat(), 'files': []}
for item in file_list:
    if item['filename'] not in selected:
        continue
    target = RAW / selected[item['filename']]
    details = item['content_details']
    if not target.exists():
        downloaded = requests.get(details['download_url'], timeout=120)
        downloaded.raise_for_status()
        target.write_bytes(downloaded.content)
    sha = hashlib.sha256(target.read_bytes()).hexdigest()
    if sha != details['sha256_hash']:
        raise ValueError(f'Source checksum does not match: {target.name}')
    manifest['files'].append({'filename': target.name, 'source_filename': item['filename'],
                              'url': details['download_url'], 'sha256': sha,
                              'bytes': target.stat().st_size})
    print('VERIFIED', target.name, target.stat().st_size, flush=True)
(ROOT / 'data/source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')

dictionary = pd.read_excel(RAW / 'variable_dictionary.xlsx')
print('VARIABLE DICTIONARY', len(dictionary), 'definitions loaded', flush=True)
wb = load_workbook(RAW / 'health_insurance_portfolio.xlsx', read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
iterator = ws.iter_rows(values_only=True)
columns = [str(x).strip() if x is not None else '' for x in next(iterator)]
print('SHEETS', wb.sheetnames, 'WORKSHEET DIMENSIONS', ws.max_row, ws.max_column, flush=True)
print('COLUMNS', columns, flush=True)
frame = pd.DataFrame(iterator, columns=columns)
wb.close()
report = {'python_version': platform.python_version(), 'pandas_version': pd.__version__,
          'rows': len(frame), 'columns': len(frame.columns), 'column_names': columns,
          'missing_values': {c: int(v) for c, v in frame.isna().sum().items()},
          'exact_duplicate_rows': int(frame.duplicated().sum()),
          'categorical_counts': {}}
for c in columns:
    if any(term in c.lower() for term in ['lapse', 'year', 'status', 'channel', 'product']):
        if frame[c].nunique(dropna=False) < 50:
            report['categorical_counts'][c] = {str(k): int(v) for k, v in frame[c].value_counts(dropna=False).items()}
for c in columns:
    if 'id' in c.lower() and frame[c].nunique() > 100:
        report.setdefault('identifier_cardinality', {})[c] = int(frame[c].nunique())
frame.to_pickle(RAW / 'portfolio_local_cache.pkl')
report['duplicate_id_period'] = int(frame.duplicated(['ID', 'period']).sum())
assert report['duplicate_id_period'] == 0, 'ID and period are not unique'
assert set(frame['lapse'].unique()) == {1, 2, 3}, 'Unexpected status codes'
report['year_counts'] = {str(k): int(v) for k, v in frame['period'].value_counts().sort_index().items()}
report['year_lapse_counts'] = {
    str(y): {str(k): int(v) for k, v in g['lapse'].value_counts().sort_index().items()}
    for y, g in frame.groupby('period')
}
report['next_year_cohorts'] = []
for year in [2017, 2018]:
    active = frame.loc[(frame['period'] == year) & (frame['lapse'] == 2), ['ID', 'ID_policy']]
    following = frame.loc[frame['period'] == year + 1, ['ID', 'lapse']]
    joined = active.merge(following, on='ID', how='left', validate='one_to_one', indicator=True)
    unknown = int(joined['lapse'].isna().sum())
    counts = {str(int(k)): int(v) for k, v in joined['lapse'].value_counts().sort_index().items()}
    report['next_year_cohorts'].append({'feature_year': year, 'outcome_year': year + 1,
        'active_at_origin': len(active), 'matched_next_year': len(joined) - unknown,
        'unknown_next_year': unknown, 'next_year_lapse_counts': counts,
        'candidate_lapse_1_or_3': counts.get('1', 0) + counts.get('3', 0)})
(ROOT / 'docs/data_feasibility.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print('FEASIBILITY', json.dumps(report, ensure_ascii=False), flush=True)
